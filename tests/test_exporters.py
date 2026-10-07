"""
tests/test_exporters.py
Unit tests for ExcelExporter and PDFExporter using a temporary database.
"""
import os
import tempfile
import pytest
from openpyxl import load_workbook
from src.db.database import Database
from src.db.session_repo import SessionRepository
from src.db.model_repo import ModelRepository
from src.db.register_library_repo import RegisterLibraryRepo
from src.db.model_map_repo import ModelMapRepo
from src.db.report_repo import ReportRepository
from src.utils.exporters import ExcelExporter, PDFExporter


@pytest.fixture
def populated_db(tmp_path):
    db_path = str(tmp_path / "test.db")
    db = Database(db_path)
    db.initialize()

    session_repo = SessionRepository(db)
    model_repo = ModelRepository(db)
    lib_repo = RegisterLibraryRepo(db)
    map_repo = ModelMapRepo(db)
    report_repo = ReportRepository(db)

    db.execute(
        "INSERT OR IGNORE INTO users (id, username, role, password_hash) VALUES (1,'test_op','OPERATOR','x')"
    )

    model_id = model_repo.create_model("MI-7646AZ", "Test Desc", "M123")

    volt_id = lib_repo.create_register(
        name="Output Voltage", register_address=100,
        register_type="HOLDING", data_type="INT16", unit="V"
    )
    curr_id = lib_repo.create_register(
        name="Output Current", register_address=101,
        register_type="HOLDING", data_type="INT16", unit="A"
    )

    map_repo.add_mapping(model_id, volt_id, role="MEASURED", display_name="Volt")
    map_repo.add_mapping(model_id, curr_id, role="MEASURED", display_name="Curr")
    params = [
        {"id": volt_id, "param_name": "Volt"},
        {"id": curr_id, "param_name": "Curr"},
    ]

    for i in range(5):
        sid = session_repo.open_session(model_id, 1)
        is_failing = (i == 4)
        for p in params:
            result = "FAIL" if is_failing else "PASS"
            session_repo.record_result(
                sid, p["id"], p["param_name"], "Test Module",
                5.0, 4.0, 6.0, result
            )
        ok = 0 if is_failing else 1
        ng = 1 if is_failing else 0
        session_repo.close_session(sid, ok, ng, 1)

    return {
        "db": db,
        "report_repo": report_repo,
        "model_id": model_id,
        "session_ids": [1, 2, 3, 4, 5]
    }


def _get_bulk_data(report_repo, model_id):
    """Helper to build bulk export data dict."""
    sessions = report_repo.get_sessions_summary(model_id)
    all_results = report_repo.get_all_results_for_range(model_id)
    return {
        "sessions": sessions,
        "results": all_results,
        "generated_at": "2026-08-06 12:00:00",
    }


class TestExcelExporter:
    def test_bulk_report_creates_file(self, tmp_path, populated_db):
        out_path = str(tmp_path / "bulk_report.xlsx")
        data = _get_bulk_data(populated_db["report_repo"], populated_db["model_id"])
        result = ExcelExporter.export_bulk(out_path, data)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0

    def test_bulk_report_has_3_sheets(self, tmp_path, populated_db):
        out_path = str(tmp_path / "bulk_report.xlsx")
        data = _get_bulk_data(populated_db["report_repo"], populated_db["model_id"])
        ExcelExporter.export_bulk(out_path, data)

        wb = load_workbook(out_path)
        assert len(wb.sheetnames) == 3
        assert "Summary" in wb.sheetnames
        assert "All Results" in wb.sheetnames
        assert "Failures" in wb.sheetnames

    def test_summary_has_5_sessions(self, tmp_path, populated_db):
        out_path = str(tmp_path / "bulk_report.xlsx")
        data = _get_bulk_data(populated_db["report_repo"], populated_db["model_id"])
        ExcelExporter.export_bulk(out_path, data)

        wb = load_workbook(out_path)
        ws = wb["Summary"]
        sessions_found = 0
        for r in range(4, 20):
            if ws.cell(row=r, column=1).value and ws.cell(row=r, column=1).value != "TOTALS":
                sessions_found += 1
        assert sessions_found == 5

    def test_session_export_creates_file(self, tmp_path, populated_db):
        out_path = str(tmp_path / "session.xlsx")
        detail = populated_db["report_repo"].get_session_detail(1)
        result = ExcelExporter.export_session(out_path, detail)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0


class TestPDFExporter:
    def test_bulk_pdf_creates_file(self, tmp_path, populated_db):
        out_path = str(tmp_path / "bulk_report.pdf")
        data = _get_bulk_data(populated_db["report_repo"], populated_db["model_id"])
        result = PDFExporter.export_bulk(out_path, data)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0

    def test_bulk_pdf_valid_header(self, tmp_path, populated_db):
        out_path = str(tmp_path / "bulk_report.pdf")
        data = _get_bulk_data(populated_db["report_repo"], populated_db["model_id"])
        PDFExporter.export_bulk(out_path, data)

        with open(out_path, 'rb') as f:
            header = f.read(4)
        assert header == b'%PDF'

    def test_session_pdf_creates_file(self, tmp_path, populated_db):
        out_path = str(tmp_path / "cert.pdf")
        detail = populated_db["report_repo"].get_session_detail(1)
        result = PDFExporter.export_session(out_path, detail)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0


class TestMappingLimits:
    def test_migration_adds_limit_columns(self, tmp_path):
        db_path = str(tmp_path / "migrate.db")
        db = Database(db_path)
        db.initialize()
        cols = {row["name"] for row in db.fetchall("PRAGMA table_info(model_register_map)")}
        assert "limit_min" in cols
        assert "limit_max" in cols

    def test_add_mapping_roundtrip_limits(self, tmp_path):
        db_path = str(tmp_path / "limits.db")
        db = Database(db_path)
        db.initialize()
        model_repo = ModelRepository(db)
        lib_repo = RegisterLibraryRepo(db)
        map_repo = ModelMapRepo(db)

        model_id = model_repo.create_model("LIM-TEST", "Desc", "L1")
        reg_id = lib_repo.create_register(
            name="Current", register_address=200,
            register_type="HOLDING", data_type="INT16", unit="A"
        )
        map_repo.add_mapping(
            model_id, reg_id, role="MEASURED", display_name="Curr",
            limit_min=1.5, limit_max=4.5
        )
        mappings = map_repo.get_model_mappings(model_id)
        assert mappings[0]["limit_min"] == 1.5
        assert mappings[0]["limit_max"] == 4.5

    def test_add_mapping_defaults_zero(self, tmp_path):
        db_path = str(tmp_path / "defaults.db")
        db = Database(db_path)
        db.initialize()
        model_repo = ModelRepository(db)
        lib_repo = RegisterLibraryRepo(db)
        map_repo = ModelMapRepo(db)

        model_id = model_repo.create_model("DEF-TEST", "Desc", "D1")
        reg_id = lib_repo.create_register(
            name="Volt", register_address=201,
            register_type="HOLDING", data_type="INT16", unit="V"
        )
        map_repo.add_mapping(model_id, reg_id, role="MEASURED", display_name="V")
        mappings = map_repo.get_model_mappings(model_id)
        assert mappings[0]["limit_min"] == 0.0
        assert mappings[0]["limit_max"] == 0.0

    def test_copy_mappings_carries_limits(self, tmp_path):
        db_path = str(tmp_path / "copy.db")
        db = Database(db_path)
        db.initialize()
        model_repo = ModelRepository(db)
        lib_repo = RegisterLibraryRepo(db)
        map_repo = ModelMapRepo(db)

        m1 = model_repo.create_model("SRC", "Desc", "S1")
        m2 = model_repo.create_model("DST", "Desc", "S2")
        reg_id = lib_repo.create_register(
            name="Curr", register_address=202,
            register_type="HOLDING", data_type="INT16", unit="A"
        )
        map_repo.add_mapping(m1, reg_id, role="MEASURED", display_name="Curr",
                             limit_min=2.0, limit_max=8.0)
        copied = map_repo.copy_mappings(m1, m2)
        assert copied == 1
        mappings = map_repo.get_model_mappings(m2)
        assert mappings[0]["limit_min"] == 2.0
        assert mappings[0]["limit_max"] == 8.0
