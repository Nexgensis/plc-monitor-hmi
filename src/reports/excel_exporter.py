"""
Excel report generation logic using openpyxl.
"""
import logging
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment

logger = logging.getLogger(__name__)

class ExcelExporter:
    NAVY_FILL = PatternFill(start_color="1E2D4A", end_color="1E2D4A", fill_type="solid")
    NAVY_FONT = Font(color="FFFFFF", bold=True)
    PASS_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    FAIL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    THIN_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    def __init__(self, report_repo):
        self.repo = report_repo

    def _safe(self, v) -> str:
        if v is None: return ""
        return str(v).encode('utf-8', 'replace').decode().strip()

    def generate_session_report(self, output_path: str, session_id: int) -> str:
        detail = self.repo.get_session_detail(session_id)
        if not detail: raise ValueError(f"Session {session_id} not found")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"
        
        ws.merge_cells("A1:F1")
        ws["A1"] = f"TEST CERTIFICATE - SESSION #{session_id}"
        ws["A1"].fill = self.NAVY_FILL
        ws["A1"].font = self.NAVY_FONT
        ws["A1"].alignment = Alignment(horizontal="center")
        
        s = detail["session"]
        ws["A3"] = "Model ID:"; ws["B3"] = s.get("model_id")
        ws["A4"] = "Start Time:"; ws["B4"] = s["started_at"]
        ws["D3"] = "OK Count:";    ws["E3"] = s["ok_count"]
        ws["D4"] = "NG Count:";    ws["E4"] = s["ng_count"]
        
        headers = ["Module", "Parameter", "Measured", "Min", "Max", "Result"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(6, col, h)
            cell.fill = self.NAVY_FILL
            cell.font = self.NAVY_FONT
            cell.border = self.THIN_BORDER
            
        row_idx = 7
        for mod, params in detail["results"].items():
            for p in params:
                ws.cell(row_idx, 1, self._safe(mod)).border = self.THIN_BORDER
                ws.cell(row_idx, 2, self._safe(p["param_name"])).border = self.THIN_BORDER
                ws.cell(row_idx, 3, p["measured_value"]).border = self.THIN_BORDER
                ws.cell(row_idx, 4, p["limit_min"]).border = self.THIN_BORDER
                ws.cell(row_idx, 5, p["limit_max"]).border = self.THIN_BORDER
                
                res_cell = ws.cell(row_idx, 6, self._safe(p["result"]))
                res_cell.border = self.THIN_BORDER
                if p["result"] == "PASS": res_cell.fill = self.PASS_FILL
                elif p["result"] == "FAIL": res_cell.fill = self.FAIL_FILL
                row_idx += 1
                
        wb.save(output_path)
        return output_path

    def generate_shift_report(self, output_path: str, model_id=None, date_from=None, date_to=None) -> str:
        sessions = self.repo.get_sessions_summary(model_id, date_from, date_to)
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        
        headers = ["ID", "Started", "Model", "OK", "NG", "Pass%"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(1, col, h)
            cell.fill = self.NAVY_FILL
            cell.font = self.NAVY_FONT
            
        for i, s in enumerate(sessions, 2):
            ws.cell(i, 1, s["session_id"])
            ws.cell(i, 2, s["started_at"])
            ws.cell(i, 3, self._safe(s["model_name"]))
            ws.cell(i, 4, s["ok_count"])
            ws.cell(i, 5, s["ng_count"])
            ws.cell(i, 6, s["pass_rate_pct"])
            
        wb.save(output_path)
        return output_path
